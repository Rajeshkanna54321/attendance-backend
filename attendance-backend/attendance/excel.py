import re
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font


def build_workbook(sessions):
    """One sheet per class session; returns xlsx bytes."""
    wb = Workbook()
    wb.remove(wb.active)
    if not sessions:
        wb.create_sheet("No sessions")
    for s in sessions:
        title = re.sub(r"[\[\]\*\?/\\:]", "", f"{s.subject[:20]}-{s.id}")
        ws = wb.create_sheet(title)
        started = timezone.localtime(s.started_at)
        ws.append([f"{s.subject} - {started:%d %b %Y %H:%M}"])
        ws["A1"].font = Font(bold=True)
        ws.append(["#", "Roll No", "Name", "Email", "Marked at", "Distance (m)"])
        for c in ws[2]:
            c.font = Font(bold=True)
        for i, r in enumerate(s.records.select_related("student").order_by("student__roll_number"), 1):
            ws.append([i, r.student.roll_number, r.student.name, r.student.email,
                       timezone.localtime(r.marked_at).strftime("%H:%M:%S"), round(r.distance_m, 1)])
        for col, w in zip("ABCDEF", (5, 14, 24, 30, 12, 13)):
            ws.column_dimensions[col].width = w
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
